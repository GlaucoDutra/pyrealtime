"""Reusable, stateless preparation of user attachments for Realtime clients."""

from __future__ import annotations

import asyncio
import base64
import io
import re
from dataclasses import asdict, dataclass, field
from pathlib import PurePath
from typing import Any, Literal

from .exceptions import AttachmentError, AttachmentTooLargeError


@dataclass(frozen=True, slots=True)
class AttachmentPolicy:
    max_file_bytes: int = 25 * 1024 * 1024
    max_text_chars: int = 120_000
    chunk_chars: int = 8_000
    max_image_dimension: int = 1_024
    max_image_data_url_bytes: int = 160_000
    max_document_pages: int = 250
    max_spreadsheet_cells: int = 100_000


@dataclass(frozen=True, slots=True)
class PreparedAttachment:
    kind: Literal["text", "image", "notice"]
    filename: str
    media_type: str
    size_bytes: int
    chunks: tuple[str, ...] = ()
    truncated: bool = False
    data_url: str | None = None
    message: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        value = asdict(self)
        value["chunks"] = list(self.chunks)
        return {key: item for key, item in value.items() if item not in (None, (), {})}


class AttachmentProcessor:
    """Convert common files into small, transport-neutral text or image payloads."""

    _text_extensions = {
        "txt", "md", "csv", "json", "jsonl", "js", "jsx", "ts", "tsx", "html",
        "css", "xml", "yaml", "yml", "toml", "py", "java", "c", "cpp", "h", "hpp",
        "sql", "sh", "ps1", "log",
    }
    _image_extensions = {"png", "jpg", "jpeg", "gif", "webp", "bmp", "tif", "tiff"}

    def __init__(self, policy: AttachmentPolicy | None = None) -> None:
        self.policy = policy or AttachmentPolicy()
        if self.policy.chunk_chars <= 0 or self.policy.max_text_chars <= 0:
            raise ValueError("Attachment text limits must be positive")

    async def prepare(self, filename: str, media_type: str, data: bytes) -> PreparedAttachment:
        return await asyncio.to_thread(self.prepare_sync, filename, media_type, data)

    def prepare_sync(self, filename: str, media_type: str, data: bytes) -> PreparedAttachment:
        safe_name = self._safe_filename(filename)
        if not data:
            raise AttachmentError("The selected file is empty")
        if len(data) > self.policy.max_file_bytes:
            raise AttachmentTooLargeError(
                f"File exceeds the {self.policy.max_file_bytes // (1024 * 1024)} MB limit"
            )

        extension = self._extension(safe_name)
        normalized_type = (media_type or "application/octet-stream").split(";", 1)[0].strip().lower()
        if extension in self._image_extensions or normalized_type.startswith("image/"):
            return self._prepare_image(safe_name, data)
        if extension in self._text_extensions or normalized_type.startswith("text/"):
            return self._prepare_text(safe_name, normalized_type, self._decode_text(data))
        if extension == "pdf" or normalized_type == "application/pdf":
            return self._prepare_text(safe_name, "application/pdf", self._extract_pdf(data))
        if extension in {"xlsx", "xlsm"}:
            return self._prepare_text(safe_name, normalized_type, self._extract_workbook(data))
        if extension == "docx":
            return self._prepare_text(safe_name, normalized_type, self._extract_docx(data))
        if extension == "pptx":
            return self._prepare_text(safe_name, normalized_type, self._extract_pptx(data))

        return PreparedAttachment(
            kind="notice",
            filename=safe_name,
            media_type=normalized_type,
            size_bytes=len(data),
            message="This file type is not supported. Convert it to PDF, plain text, DOCX, PPTX, or XLSX first.",
        )

    def _prepare_text(self, filename: str, media_type: str, text: str) -> PreparedAttachment:
        clipped = text[: self.policy.max_text_chars]
        chunks = tuple(
            clipped[offset : offset + self.policy.chunk_chars]
            for offset in range(0, len(clipped), self.policy.chunk_chars)
        ) or ("[empty file]",)
        return PreparedAttachment(
            kind="text",
            filename=filename,
            media_type=media_type,
            size_bytes=len(text.encode("utf-8")),
            chunks=chunks,
            truncated=len(text) > self.policy.max_text_chars,
            metadata={"character_count": len(text)},
        )

    def _prepare_image(self, filename: str, data: bytes) -> PreparedAttachment:
        try:
            from PIL import Image, ImageOps
        except ImportError as exc:  # pragma: no cover - installation guidance
            raise AttachmentError("Image support requires pyrealtime[files]") from exc
        try:
            with Image.open(io.BytesIO(data)) as source:
                image = ImageOps.exif_transpose(source).convert("RGB")
                image.thumbnail((self.policy.max_image_dimension, self.policy.max_image_dimension))
                quality = 78
                encoded = b""
                for _ in range(8):
                    output = io.BytesIO()
                    image.save(output, format="JPEG", quality=quality, optimize=True)
                    encoded = output.getvalue()
                    data_url_bytes = 23 + 4 * ((len(encoded) + 2) // 3)
                    if data_url_bytes <= self.policy.max_image_data_url_bytes:
                        break
                    quality = max(38, quality - 8)
                    image.thumbnail((max(320, int(image.width * 0.82)), max(320, int(image.height * 0.82))))
                data_url_bytes = 23 + 4 * ((len(encoded) + 2) // 3)
                if data_url_bytes > self.policy.max_image_data_url_bytes:
                    raise AttachmentTooLargeError("The normalized image is too large for a Realtime message")
                width, height = image.size
        except AttachmentError:
            raise
        except Exception as exc:
            raise AttachmentError("The image could not be decoded") from exc
        data_url = "data:image/jpeg;base64," + base64.b64encode(encoded).decode("ascii")
        return PreparedAttachment(
            kind="image",
            filename=filename,
            media_type="image/jpeg",
            size_bytes=len(data),
            data_url=data_url,
            metadata={"width": width, "height": height, "normalized_bytes": len(encoded)},
        )

    def _extract_pdf(self, data: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover
            raise AttachmentError("PDF support requires pyrealtime[files]") from exc
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise AttachmentError("Encrypted PDFs are not supported")
            if len(reader.pages) > self.policy.max_document_pages:
                raise AttachmentError(f"PDF exceeds the {self.policy.max_document_pages}-page limit")
            pages = []
            for index, page in enumerate(reader.pages, 1):
                text = (page.extract_text() or "").strip()
                pages.append(f"[PAGE {index}]\n{text or '[no selectable text found on this page]'}")
            return "\n\n".join(pages)
        except AttachmentError:
            raise
        except Exception as exc:
            raise AttachmentError("The PDF could not be read") from exc

    def _extract_workbook(self, data: bytes) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise AttachmentError("Spreadsheet support requires pyrealtime[files]") from exc
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines: list[str] = []
            cells = 0
            for sheet in workbook.worksheets:
                lines.append(f"[SHEET {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    cells += len(row)
                    if cells > self.policy.max_spreadsheet_cells:
                        raise AttachmentError("Spreadsheet exceeds the cell-processing limit")
                    values = [
                        re.sub(r"[\t\r\n]+", " ", "" if value is None else str(value)).strip()
                        for value in row
                    ]
                    if any(values):
                        lines.append("\t".join(values).rstrip())
            workbook.close()
            return "\n".join(lines) or "[empty workbook]"
        except AttachmentError:
            raise
        except Exception as exc:
            raise AttachmentError("The spreadsheet could not be read") from exc

    def _extract_docx(self, data: bytes) -> str:
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover
            raise AttachmentError("DOCX support requires pyrealtime[files]") from exc
        try:
            document = Document(io.BytesIO(data))
            lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table_number, table in enumerate(document.tables, 1):
                lines.append(f"[TABLE {table_number}]")
                for row in table.rows:
                    lines.append("\t".join(cell.text.replace("\n", " ").strip() for cell in row.cells))
            return "\n".join(lines) or "[empty document]"
        except Exception as exc:
            raise AttachmentError("The DOCX document could not be read") from exc

    def _extract_pptx(self, data: bytes) -> str:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            raise AttachmentError("PPTX support requires pyrealtime[files]") from exc
        try:
            presentation = Presentation(io.BytesIO(data))
            if len(presentation.slides) > self.policy.max_document_pages:
                raise AttachmentError(f"Presentation exceeds the {self.policy.max_document_pages}-slide limit")
            slides: list[str] = []
            for index, slide in enumerate(presentation.slides, 1):
                text = "\n".join(
                    str(shape.text).strip() for shape in slide.shapes
                    if hasattr(shape, "text") and str(shape.text).strip()
                )
                slides.append(f"[SLIDE {index}]\n{text or '[no text found on this slide]'}")
            return "\n\n".join(slides) or "[empty presentation]"
        except AttachmentError:
            raise
        except Exception as exc:
            raise AttachmentError("The PPTX presentation could not be read") from exc

    @staticmethod
    def _decode_text(data: bytes) -> str:
        if b"\x00" in data[:4096]:
            raise AttachmentError("The file appears to be binary, not text")
        try:
            return data.decode("utf-8-sig")
        except UnicodeDecodeError:
            return data.decode("utf-8", errors="replace")

    @staticmethod
    def _safe_filename(filename: str) -> str:
        clean = PurePath((filename or "").replace("\\", "/")).name
        clean = re.sub(r"[\x00-\x1f\x7f]", "", clean).strip()[:255]
        if not clean or clean in {".", ".."}:
            raise AttachmentError("A valid file name is required")
        return clean

    @staticmethod
    def _extension(filename: str) -> str:
        return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
